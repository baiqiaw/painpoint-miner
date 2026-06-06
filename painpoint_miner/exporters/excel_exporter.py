"""Excel 导出（含格式化）。"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..models.base import AnalysisReport
from .base import BaseExporter


class ExcelExporter(BaseExporter):
    """Excel 导出器。生成格式化的 xlsx 文件。"""

    HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)

    @property
    def format_name(self) -> str:
        return "excel"

    async def export(self, report: AnalysisReport, output_path: Path) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()

        # Sheet 1: 痛点概览
        ws_pp = wb.active
        ws_pp.title = "痛点概览"
        self._write_pain_points(ws_pp, report)

        # Sheet 2: 主题聚类
        ws_cl = wb.create_sheet("主题聚类")
        self._write_clusters(ws_cl, report)

        # Sheet 3: 平台统计
        ws_pl = wb.create_sheet("平台统计")
        self._write_platform_stats(ws_pl, report)

        wb.save(output_path)
        return output_path

    def _style_header(self, ws, headers: list, row: int = 1):
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

    def _write_pain_points(self, ws, report: AnalysisReport):
        headers = ["#", "描述", "类型", "严重度", "情感分数", "情感标签", "频率", "平台", "引用"]
        self._style_header(ws, headers)

        for i, pp in enumerate(report.pain_points, 1):
            ws.cell(row=i + 1, column=1, value=i)
            ws.cell(row=i + 1, column=2, value=pp.description)
            ws.cell(row=i + 1, column=3, value=pp.pain_type.value)
            ws.cell(row=i + 1, column=4, value=pp.severity.value)
            ws.cell(row=i + 1, column=5, value=pp.sentiment_score)
            ws.cell(row=i + 1, column=6, value=pp.sentiment_label.value)
            ws.cell(row=i + 1, column=7, value=pp.frequency)
            ws.cell(row=i + 1, column=8, value=", ".join(p.value for p in pp.platforms))
            ws.cell(row=i + 1, column=9, value="; ".join(pp.evidence_quotes))

        # 自动列宽
        for col in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)

    def _write_clusters(self, ws, report: AnalysisReport):
        headers = ["簇ID", "标签", "大小", "关键词", "痛点数"]
        self._style_header(ws, headers)

        for i, cluster in enumerate(report.topic_clusters, 1):
            ws.cell(row=i + 1, column=1, value=cluster.id)
            ws.cell(row=i + 1, column=2, value=cluster.label)
            ws.cell(row=i + 1, column=3, value=cluster.size)
            ws.cell(row=i + 1, column=4, value=", ".join(cluster.keywords))
            ws.cell(row=i + 1, column=5, value=len(cluster.pain_point_ids))

    def _write_platform_stats(self, ws, report: AnalysisReport):
        headers = ["平台", "帖子数"]
        self._style_header(ws, headers)

        for i, (platform, count) in enumerate(report.platform_breakdown.items(), 1):
            ws.cell(row=i + 1, column=1, value=platform)
            ws.cell(row=i + 1, column=2, value=count)
