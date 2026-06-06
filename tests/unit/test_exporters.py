"""导出器单元测试。"""

import json
from pathlib import Path

import pytest

from painpoint_miner.exporters import BaseExporter, ExcelExporter, JsonExporter, MarkdownExporter
from painpoint_miner.models.base import (
    AnalysisReport,
    CostSummary,
    PainPoint,
    TopicCluster,
)
from painpoint_miner.models.enums import (
    PainType,
    Platform,
    Sentiment,
    Severity,
)


# ── fixtures ──────────────────────────────────────────────────────


@pytest.fixture
def pain_points():
    """两个示例痛点。"""
    return [
        PainPoint(
            id="pp_001",
            source_post_ids=["xhs_123", "wb_456"],
            platforms=[Platform.XIAOHONGSHU, Platform.WEIBO],
            description="按钮位置不明显，用户难以找到功能入口",
            pain_type=PainType.UX_PROBLEM,
            severity=Severity.MAJOR,
            sentiment_score=-0.7,
            sentiment_label=Sentiment.NEGATIVE,
            frequency=15,
            evidence_quotes=["按钮找不到在哪里", "设计太差了"],
        ),
        PainPoint(
            id="pp_002",
            source_post_ids=["dy_789"],
            platforms=[Platform.DOUYIN],
            description="加载速度太慢，经常卡顿",
            pain_type=PainType.PERFORMANCE,
            severity=Severity.CRITICAL,
            sentiment_score=-0.9,
            sentiment_label=Sentiment.VERY_NEGATIVE,
            frequency=23,
            evidence_quotes=["一直转圈圈", "等半天加载不出来"],
        ),
    ]


@pytest.fixture
def topic_clusters():
    """两个示例聚类。"""
    return [
        TopicCluster(
            id=0,
            label="UI/UX 设计问题",
            pain_point_ids=["pp_001"],
            keywords=["按钮", "界面", "找不到"],
            size=1,
        ),
        TopicCluster(
            id=1,
            label="性能问题",
            pain_point_ids=["pp_002"],
            keywords=["加载", "卡顿", "慢"],
            size=1,
        ),
    ]


@pytest.fixture
def report(pain_points, topic_clusters):
    """完整分析报告。"""
    return AnalysisReport(
        query_keywords=["产品难用", "太贵了"],
        total_posts_scraped=200,
        total_comments_scraped=500,
        pain_points=pain_points,
        topic_clusters=topic_clusters,
        summary="用户主要反馈集中在 UI 设计和性能两方面。",
        platform_breakdown={"xiaohongshu": 80, "weibo": 60, "douyin": 40, "zhihu": 20},
        cost_summary=CostSummary(
            total_input_tokens=10000,
            total_output_tokens=2000,
            total_calls=5,
            estimated_cost_usd=0.15,
            model_used="claude-haiku-4-5-20251001",
        ),
    )


@pytest.fixture
def empty_report():
    """空报告（边界条件）。"""
    return AnalysisReport(
        query_keywords=["测试"],
        total_posts_scraped=0,
        pain_points=[],
        topic_clusters=[],
        summary="",
        platform_breakdown={},
    )


# ── BaseExporter 抽象 ────────────────────────────────────────────


class TestBaseExporter:
    """BaseExporter ABC 测试。"""

    def test_cannot_instantiate(self):
        """抽象类不能直接实例化。"""
        with pytest.raises(TypeError):
            BaseExporter()

    def test_subclass_must_implement(self):
        """子类必须实现所有抽象方法。"""

        class Incomplete(BaseExporter):
            pass

        with pytest.raises(TypeError):
            Incomplete()

    def test_complete_subclass(self):
        """完整子类可以实例化。"""

        class Complete(BaseExporter):
            @property
            def format_name(self) -> str:
                return "test"

            async def export(self, report, output_path):
                return output_path

        exporter = Complete()
        assert exporter.format_name == "test"


# ── JsonExporter ──────────────────────────────────────────────────


class TestJsonExporter:
    """JSON 导出器测试。"""

    def test_format_name(self):
        assert JsonExporter().format_name == "json"

    @pytest.mark.asyncio
    async def test_export_creates_file(self, report, tmp_path):
        exporter = JsonExporter()
        output = tmp_path / "report.json"
        result = await exporter.export(report, output)

        assert result == output
        assert output.exists()

    @pytest.mark.asyncio
    async def test_export_valid_json(self, report, tmp_path):
        exporter = JsonExporter()
        output = tmp_path / "report.json"
        await exporter.export(report, output)

        with open(output, encoding="utf-8") as f:
            data = json.load(f)

        assert isinstance(data, dict)

    @pytest.mark.asyncio
    async def test_export_contains_keywords(self, report, tmp_path):
        exporter = JsonExporter()
        output = tmp_path / "report.json"
        await exporter.export(report, output)

        with open(output, encoding="utf-8") as f:
            data = json.load(f)

        assert data["query_keywords"] == ["产品难用", "太贵了"]

    @pytest.mark.asyncio
    async def test_export_contains_pain_points(self, report, tmp_path):
        exporter = JsonExporter()
        output = tmp_path / "report.json"
        await exporter.export(report, output)

        with open(output, encoding="utf-8") as f:
            data = json.load(f)

        assert len(data["pain_points"]) == 2
        pp = data["pain_points"][0]
        assert pp["description"] == "按钮位置不明显，用户难以找到功能入口"
        assert pp["pain_type"] == "ux_problem"
        assert pp["severity"] == 4  # MAJOR = 4

    @pytest.mark.asyncio
    async def test_export_contains_clusters(self, report, tmp_path):
        exporter = JsonExporter()
        output = tmp_path / "report.json"
        await exporter.export(report, output)

        with open(output, encoding="utf-8") as f:
            data = json.load(f)

        assert len(data["topic_clusters"]) == 2
        assert data["topic_clusters"][0]["label"] == "UI/UX 设计问题"

    @pytest.mark.asyncio
    async def test_export_contains_platform_breakdown(self, report, tmp_path):
        exporter = JsonExporter()
        output = tmp_path / "report.json"
        await exporter.export(report, output)

        with open(output, encoding="utf-8") as f:
            data = json.load(f)

        assert data["platform_breakdown"]["xiaohongshu"] == 80

    @pytest.mark.asyncio
    async def test_export_contains_cost_summary(self, report, tmp_path):
        exporter = JsonExporter()
        output = tmp_path / "report.json"
        await exporter.export(report, output)

        with open(output, encoding="utf-8") as f:
            data = json.load(f)

        assert data["cost_summary"]["estimated_cost_usd"] == 0.15

    @pytest.mark.asyncio
    async def test_export_empty_report(self, empty_report, tmp_path):
        exporter = JsonExporter()
        output = tmp_path / "report.json"
        await exporter.export(empty_report, output)

        with open(output, encoding="utf-8") as f:
            data = json.load(f)

        assert data["pain_points"] == []
        assert data["topic_clusters"] == []
        assert data["platform_breakdown"] == {}

    @pytest.mark.asyncio
    async def test_export_creates_parent_dirs(self, report, tmp_path):
        exporter = JsonExporter()
        output = tmp_path / "sub" / "dir" / "report.json"
        result = await exporter.export(report, output)

        assert result.exists()
        assert output.parent.exists()

    @pytest.mark.asyncio
    async def test_export_chinese_not_escaped(self, report, tmp_path):
        """确保中文字符不被 Unicode 转义。"""
        exporter = JsonExporter()
        output = tmp_path / "report.json"
        await exporter.export(report, output)

        raw = output.read_text(encoding="utf-8")
        assert "产品难用" in raw  # 未被 \uXXXX 转义


# ── MarkdownExporter ──────────────────────────────────────────────


class TestMarkdownExporter:
    """Markdown 报告导出器测试。"""

    def test_format_name(self):
        assert MarkdownExporter().format_name == "markdown"

    @pytest.mark.asyncio
    async def test_export_creates_file(self, report, tmp_path):
        exporter = MarkdownExporter()
        output = tmp_path / "report.md"
        result = await exporter.export(report, output)

        assert result == output
        assert output.exists()

    @pytest.mark.asyncio
    async def test_export_contains_title(self, report, tmp_path):
        exporter = MarkdownExporter()
        output = tmp_path / "report.md"
        await exporter.export(report, output)

        content = output.read_text(encoding="utf-8")
        assert "# 用户痛点分析报告" in content

    @pytest.mark.asyncio
    async def test_export_contains_keywords(self, report, tmp_path):
        exporter = MarkdownExporter()
        output = tmp_path / "report.md"
        await exporter.export(report, output)

        content = output.read_text(encoding="utf-8")
        assert "产品难用" in content
        assert "太贵了" in content

    @pytest.mark.asyncio
    async def test_export_contains_platform_table(self, report, tmp_path):
        exporter = MarkdownExporter()
        output = tmp_path / "report.md"
        await exporter.export(report, output)

        content = output.read_text(encoding="utf-8")
        assert "| 平台 | 帖子数 |" in content
        assert "| xiaohongshu | 80 |" in content

    @pytest.mark.asyncio
    async def test_export_contains_clusters(self, report, tmp_path):
        exporter = MarkdownExporter()
        output = tmp_path / "report.md"
        await exporter.export(report, output)

        content = output.read_text(encoding="utf-8")
        assert "UI/UX 设计问题" in content
        assert "性能问题" in content

    @pytest.mark.asyncio
    async def test_export_contains_pain_point_details(self, report, tmp_path):
        exporter = MarkdownExporter()
        output = tmp_path / "report.md"
        await exporter.export(report, output)

        content = output.read_text(encoding="utf-8")
        assert "按钮位置不明显" in content
        assert "加载速度太慢" in content
        assert "ux_problem" in content
        assert "performance" in content

    @pytest.mark.asyncio
    async def test_export_contains_severity_stars(self, report, tmp_path):
        """严重度用星级表示。"""
        exporter = MarkdownExporter()
        output = tmp_path / "report.md"
        await exporter.export(report, output)

        content = output.read_text(encoding="utf-8")
        # MAJOR=4 → ⭐⭐⭐⭐
        assert "⭐⭐⭐⭐ (4/5)" in content
        # CRITICAL=5 → ⭐⭐⭐⭐⭐
        assert "⭐⭐⭐⭐⭐ (5/5)" in content

    @pytest.mark.asyncio
    async def test_export_contains_evidence_quotes(self, report, tmp_path):
        exporter = MarkdownExporter()
        output = tmp_path / "report.md"
        await exporter.export(report, output)

        content = output.read_text(encoding="utf-8")
        assert "> \"按钮找不到在哪里\"" in content

    @pytest.mark.asyncio
    async def test_export_contains_summary(self, report, tmp_path):
        exporter = MarkdownExporter()
        output = tmp_path / "report.md"
        await exporter.export(report, output)

        content = output.read_text(encoding="utf-8")
        assert "用户主要反馈集中在 UI 设计和性能两方面" in content

    @pytest.mark.asyncio
    async def test_export_empty_report(self, empty_report, tmp_path):
        exporter = MarkdownExporter()
        output = tmp_path / "report.md"
        await exporter.export(empty_report, output)

        content = output.read_text(encoding="utf-8")
        assert "# 用户痛点分析报告" in content
        assert "未发现痛点" in content
        assert "无数据" in content

    @pytest.mark.asyncio
    async def test_export_creates_parent_dirs(self, report, tmp_path):
        exporter = MarkdownExporter()
        output = tmp_path / "deep" / "dir" / "report.md"
        result = await exporter.export(report, output)

        assert result.exists()

    @pytest.mark.asyncio
    async def test_export_contains_disclaimer(self, report, tmp_path):
        exporter = MarkdownExporter()
        output = tmp_path / "report.md"
        await exporter.export(report, output)

        content = output.read_text(encoding="utf-8")
        assert "TikHub API" in content
        assert "responsibility" in content.lower() or "assumes" in content.lower()


# ── ExcelExporter ─────────────────────────────────────────────────


class TestExcelExporter:
    """Excel 导出器测试。"""

    def test_format_name(self):
        assert ExcelExporter().format_name == "excel"

    @pytest.mark.asyncio
    async def test_export_creates_file(self, report, tmp_path):
        exporter = ExcelExporter()
        output = tmp_path / "report.xlsx"
        result = await exporter.export(report, output)

        assert result == output
        assert output.exists()

    @pytest.mark.asyncio
    async def test_export_has_three_sheets(self, report, tmp_path):
        exporter = ExcelExporter()
        output = tmp_path / "report.xlsx"
        await exporter.export(report, output)

        from openpyxl import load_workbook

        wb = load_workbook(output)
        sheet_names = wb.sheetnames
        assert "痛点概览" in sheet_names
        assert "主题聚类" in sheet_names
        assert "平台统计" in sheet_names
        assert len(sheet_names) == 3

    @pytest.mark.asyncio
    async def test_pain_points_sheet_data(self, report, tmp_path):
        exporter = ExcelExporter()
        output = tmp_path / "report.xlsx"
        await exporter.export(report, output)

        from openpyxl import load_workbook

        wb = load_workbook(output)
        ws = wb["痛点概览"]

        # Header row
        headers = [cell.value for cell in ws[1]]
        assert "#" in headers
        assert "描述" in headers
        assert "严重度" in headers

        # Data rows (2 pain points)
        assert ws.cell(row=2, column=2).value == "按钮位置不明显，用户难以找到功能入口"
        assert ws.cell(row=2, column=4).value == 4  # MAJOR
        assert ws.cell(row=3, column=2).value == "加载速度太慢，经常卡顿"
        assert ws.cell(row=3, column=4).value == 5  # CRITICAL

    @pytest.mark.asyncio
    async def test_clusters_sheet_data(self, report, tmp_path):
        exporter = ExcelExporter()
        output = tmp_path / "report.xlsx"
        await exporter.export(report, output)

        from openpyxl import load_workbook

        wb = load_workbook(output)
        ws = wb["主题聚类"]

        # Header
        headers = [cell.value for cell in ws[1]]
        assert "簇ID" in headers
        assert "标签" in headers

        # Data
        assert ws.cell(row=2, column=2).value == "UI/UX 设计问题"
        assert ws.cell(row=3, column=2).value == "性能问题"

    @pytest.mark.asyncio
    async def test_platform_stats_sheet_data(self, report, tmp_path):
        exporter = ExcelExporter()
        output = tmp_path / "report.xlsx"
        await exporter.export(report, output)

        from openpyxl import load_workbook

        wb = load_workbook(output)
        ws = wb["平台统计"]

        # Header
        headers = [cell.value for cell in ws[1]]
        assert "平台" in headers
        assert "帖子数" in headers

        # 4 platforms
        platform_cells = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).value for i in range(2, 6)}
        assert platform_cells.get("xiaohongshu") == 80
        assert platform_cells.get("weibo") == 60

    @pytest.mark.asyncio
    async def test_header_styling(self, report, tmp_path):
        """表头应有蓝色背景和白色字体。"""
        exporter = ExcelExporter()
        output = tmp_path / "report.xlsx"
        await exporter.export(report, output)

        from openpyxl import load_workbook

        wb = load_workbook(output)
        ws = wb["痛点概览"]
        header_cell = ws.cell(row=1, column=1)

        assert header_cell.font.bold is True
        assert header_cell.font.color.rgb == "00FFFFFF"
        assert header_cell.fill.start_color.rgb == "004472C4"

    @pytest.mark.asyncio
    async def test_export_empty_report(self, empty_report, tmp_path):
        exporter = ExcelExporter()
        output = tmp_path / "report.xlsx"
        await exporter.export(empty_report, output)

        from openpyxl import load_workbook

        wb = load_workbook(output)
        ws = wb["痛点概览"]
        # Only header row, no data
        assert ws.cell(row=2, column=1).value is None

    @pytest.mark.asyncio
    async def test_export_creates_parent_dirs(self, report, tmp_path):
        exporter = ExcelExporter()
        output = tmp_path / "nested" / "path" / "report.xlsx"
        result = await exporter.export(report, output)

        assert result.exists()

    @pytest.mark.asyncio
    async def test_sentiment_score_written(self, report, tmp_path):
        exporter = ExcelExporter()
        output = tmp_path / "report.xlsx"
        await exporter.export(report, output)

        from openpyxl import load_workbook

        wb = load_workbook(output)
        ws = wb["痛点概览"]

        assert ws.cell(row=2, column=5).value == -0.7
        assert ws.cell(row=3, column=5).value == -0.9
